"""Zona del profesor: material para clase y edición del vocabulario del curso."""
from __future__ import annotations

import base64
import hmac
import io
import zipfile
from pathlib import Path
from urllib.parse import quote

import requests
import streamlit as st
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

import core

ARCHIVO_VOCAB = "vocabulario.txt"
ARCHIVO_FRASES = "training set.txt"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TIPOS_ESCRITURA = ["es_zh", "ordenar", "completar", "responder", "redaccion"]
ETIQUETAS = {
    "es_zh": "Traducir al chino", "ordenar": "Ordenar palabras", "completar": "Completar espacios",
    "responder": "Responder preguntas", "redaccion": "Redacción con palabras dadas",
}


# ---------------------------------------------------------------------------
# Documentos Word
# ---------------------------------------------------------------------------
def _nuevo_doc(titulo: str, subtitulo: str, estudiante: bool) -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.add_heading(titulo, level=1)
    doc.add_paragraph(subtitulo)
    if estudiante:
        doc.add_paragraph("Nombre: ______________________________     Fecha: ______________")
    return doc


def _chino(doc, texto: str, tam: int = 16):
    run = doc.add_paragraph().add_run(texto)
    run.font.size = Pt(tam)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def _respuesta(doc, texto: str | None, lineas: int = 1):
    """Línea para escribir (hoja del estudiante) o la respuesta en rojo (clave)."""
    if texto is None:
        for _ in range(lineas):
            doc.add_paragraph("_" * 62)
    else:
        run = doc.add_paragraph().add_run(f"Respuesta: {texto}")
        run.font.color.rgb, run.bold = RGBColor(0xC8, 0x10, 0x2E), True


def doc_examen(titulo: str, subtitulo: str, preguntas: list[dict], clave: bool) -> bytes:
    doc = _nuevo_doc(titulo + (" — Clave de respuestas" if clave else ""), subtitulo, not clave)
    for i, p in enumerate(preguntas, 1):
        doc.add_paragraph().add_run(f"{i}. {core.TIPOS[p['tipo']]}").bold = True
        tipo = p["tipo"]
        if tipo == "ordenar":
            _chino(doc, "      ".join(p["palabras"]))
            _respuesta(doc, p["respuesta"] if clave else None)
        elif tipo == "redaccion":
            _chino(doc, "，  ".join(p["palabras"]))
            doc.add_paragraph(f"Escribe un texto de unos {p['longitud']} caracteres.")
            _respuesta(doc, f"(texto modelo) {p['respuesta']}" if clave else None, lineas=4)
        elif tipo == "lectura":
            _chino(doc, p["texto"], 14)
            for j, s in enumerate(p["preguntas"], 1):
                _chino(doc, f"{chr(96 + j)}) {s['enunciado']}", 13)
                _respuesta(doc, s["respuesta"] if clave else None)
        else:
            if tipo == "es_zh":
                doc.add_paragraph(p["enunciado"])
            else:
                _chino(doc, p["enunciado"])
            _respuesta(doc, p["respuesta"] if clave else None)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def doc_listening(titulo: str, subtitulo: str, frases: list[dict], dictado: bool, clave: bool) -> bytes:
    doc = _nuevo_doc(titulo + (" — Clave de respuestas" if clave else ""), subtitulo, not clave)
    doc.add_paragraph(
        "Escucha cada audio y escribe la frase en chino (hanzi o pinyin)." if dictado
        else "Escucha cada audio y escribe en español qué significa."
    )
    for i, f in enumerate(frases, 1):
        doc.add_paragraph().add_run(f"{i}.  Audio {i:02d}").bold = True
        if clave:
            _chino(doc, f["zh"])
            doc.add_paragraph(f"{f['pinyin']}  —  {f['es']}")
        else:
            _respuesta(doc, None, lineas=1 if dictado else 2)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Guardar en GitHub
# ---------------------------------------------------------------------------
def subir_a_github(token: str, repo: str, rama: str, archivos: dict[str, str], mensaje: str) -> int:
    """Crea o actualiza los archivos en el repositorio. Devuelve cuántos cambió."""
    cab = {"Authorization": f"Bearer {token}", "Accept": "application/vnd.github+json"}
    cambiados = 0
    for ruta, texto in archivos.items():
        url = f"https://api.github.com/repos/{repo}/contents/{quote(ruta)}"
        r = requests.get(url, headers=cab, params={"ref": rama}, timeout=20)
        if r.status_code not in (200, 404):
            r.raise_for_status()
        sha = r.json().get("sha") if r.status_code == 200 else None
        if sha and base64.b64decode(r.json().get("content", "")).decode("utf-8", "replace") == texto:
            continue
        cuerpo = {"message": mensaje, "branch": rama, "content": base64.b64encode(texto.encode("utf-8")).decode()}
        if sha:
            cuerpo["sha"] = sha
        requests.put(url, headers=cab, json=cuerpo, timeout=20).raise_for_status()
        cambiados += 1
    return cambiados


# ---------------------------------------------------------------------------
# Interfaz
# ---------------------------------------------------------------------------
def _acceso(clave_correcta: str) -> bool:
    S = st.session_state
    if S.get("profe_ok"):
        return True
    if S.get("profe_intentos", 0) >= 5:
        st.error("Demasiados intentos. Recarga la página para volver a intentarlo.")
        return False
    with st.form("acceso_profe"):
        intento = st.text_input("Clave de profesor", type="password")
        entrar = st.form_submit_button("Entrar", type="primary")
    if entrar:
        if hmac.compare_digest(intento.strip().encode(), str(clave_correcta).encode()):
            S.profe_ok = True
            st.rerun()
        S.profe_intentos = S.get("profe_intentos", 0) + 1
        st.error("Clave incorrecta.")
    return False


def _material_para(vocab, frases) -> tuple[core.Contexto, str]:
    a, b = st.columns([1, 3])
    nivel = a.selectbox("Nivel", sorted(vocab), format_func=lambda n: f"Chino {n}", key="p_nivel")
    unidades = sorted(vocab[nivel])
    elegidas = b.multiselect(
        "Vocabularios", unidades, default=unidades[-1:], format_func=lambda u: f"Vocabulario {u}", key=f"p_unid_{nivel}"
    ) or unidades[-1:]
    desc = f"Chino {nivel} · Vocabulario {', '.join(str(u) for u in sorted(elegidas))}"
    return core.construir_contexto(vocab, frases, nivel, elegidas), desc


def _vista_pregunta(i: int, p: dict):
    st.markdown(f"**{i}. {core.TIPOS[p['tipo']]}**")
    if p["tipo"] == "lectura":
        st.markdown(p["texto"])
        for s in p["preguntas"]:
            st.markdown(f"- {s['enunciado']} → **{s['respuesta']}**")
    elif p["tipo"] in ("ordenar", "redaccion"):
        st.markdown(" / ".join(p["palabras"]) + f" → **{p['respuesta']}**")
    else:
        st.markdown(f"{p['enunciado']} → **{p['respuesta']}**")


def _mostrar_material(clave_estado: str, nombre_archivo: str):
    """Vista previa con respuestas (solo profesor) y descargas en Word."""
    mat = st.session_state.get(clave_estado)
    if not mat:
        return
    st.success(f"{len(mat['preguntas'])} preguntas listas. Las respuestas solo se ven aquí y en la clave.")
    a, b = st.columns(2)
    a.download_button(
        "⬇ Hoja del estudiante (Word)", doc_examen(mat["titulo"], mat["desc"], mat["preguntas"], False),
        f"{nombre_archivo}.docx", MIME_DOCX, use_container_width=True, key=f"{clave_estado}_d1",
    )
    b.download_button(
        "⬇ Clave de respuestas (Word)", doc_examen(mat["titulo"], mat["desc"], mat["preguntas"], True),
        f"{nombre_archivo}_clave.docx", MIME_DOCX, use_container_width=True, key=f"{clave_estado}_d2",
    )
    with st.container(border=True):
        for i, p in enumerate(mat["preguntas"], 1):
            _vista_pregunta(i, p)


def _generar(tutor, ctx, desc, clave_estado, titulo, dificultad, n, plan=None):
    with st.spinner("Generando…"):
        try:
            preguntas = tutor.generar_examen(ctx, dificultad, n, plan)
        except Exception as err:
            st.error(f"No se pudo generar en este momento ({type(err).__name__}). Intenta de nuevo.")
            return
    st.session_state[clave_estado] = {"titulo": titulo, "desc": desc, "preguntas": preguntas}


def _tab_examen(tutor, ctx, desc):
    st.caption("Genera un examen con su clave de respuestas para imprimir o editar en Word.")
    titulo = st.text_input("Título", "Examen de práctica", key="pe_titulo")
    a, b = st.columns(2)
    dificultad = a.selectbox("Dificultad", list(core.DIFICULTADES), index=1, key="pe_dif")
    n = b.select_slider("Número de preguntas", [5, 8, 10, 12, 15], value=10, key="pe_n")
    st.caption(core.DIFICULTADES[dificultad]["ayuda"])
    if st.button("Generar examen", type="primary", key="pe_go"):
        _generar(tutor, ctx, desc, "prof_examen", titulo, dificultad, n)
    _mostrar_material("prof_examen", "examen")


def _tab_escritura(tutor, ctx, desc):
    st.caption("Ejercicios para que el estudiante escriba en chino.")
    titulo = st.text_input("Título", "Ejercicios de escritura", key="pw_titulo")
    tipos = st.multiselect(
        "Tipos de ejercicio", TIPOS_ESCRITURA, default=["es_zh", "ordenar", "redaccion"],
        format_func=ETIQUETAS.get, key="pw_tipos",
    )
    a, b = st.columns(2)
    dificultad = a.selectbox("Dificultad", list(core.DIFICULTADES), index=1, key="pw_dif")
    n = b.select_slider("Número de ejercicios", [4, 6, 8, 10, 12], value=8, key="pw_n")
    if st.button("Generar ejercicios", type="primary", disabled=not tipos, key="pw_go"):
        cortos = [t for t in tipos if t != "redaccion"]
        if "redaccion" not in tipos:
            redacciones = 0
        else:
            redacciones = max(1, min(2, n // 4)) if cortos else min(n, 3)
        plan = [cortos[i % len(cortos)] for i in range(n - redacciones)] if cortos else []
        _generar(tutor, ctx, desc, "prof_escritura", titulo, dificultad, n, plan + ["redaccion"] * redacciones)
    _mostrar_material("prof_escritura", "escritura")


def _tab_listening(tutor, ctx, desc, sintetizar):
    st.caption("Genera frases con su audio en MP3, la hoja del estudiante y la clave.")
    titulo = st.text_input("Título", "Ejercicio de listening", key="pl_titulo")
    a, b, c, d = st.columns(4)
    n = a.select_slider("Audios", [3, 5, 8, 10], value=5, key="pl_n")
    dictado = b.selectbox("Ejercicio", ["Dictado", "Comprensión"], key="pl_tipo") == "Dictado"
    lento = c.selectbox("Velocidad", ["Lenta", "Normal"], key="pl_vel") == "Lenta"
    voz = core.VOCES[d.selectbox("Voz", list(core.VOCES), key="pl_voz")]
    if st.button("Generar listening", type="primary", key="pl_go"):
        with st.spinner("Generando frases y audios…"):
            try:
                frases = tutor.generar_frases(ctx, n)
            except Exception as err:
                st.error(f"No se pudo generar en este momento ({type(err).__name__}). Intenta de nuevo.")
                return
            audios = []
            for f in frases:
                try:
                    audios.append(sintetizar(f["zh"], voz, lento))
                except Exception:
                    audios.append(None)
        st.session_state.prof_listening = {
            "titulo": titulo, "desc": desc, "frases": frases, "audios": audios, "dictado": dictado,
        }
    mat = st.session_state.get("prof_listening")
    if not mat:
        return
    if any(a is None for a in mat["audios"]):
        st.warning("Algunos audios no se pudieron generar. Vuelve a generar para reintentarlo.")
    paquete = io.BytesIO()
    with zipfile.ZipFile(paquete, "w", zipfile.ZIP_DEFLATED) as z:
        for i, audio in enumerate(mat["audios"], 1):
            if audio:
                z.writestr(f"audio_{i:02d}.mp3", audio)
        z.writestr("hoja_estudiante.docx", doc_listening(mat["titulo"], mat["desc"], mat["frases"], mat["dictado"], False))
        z.writestr("clave_respuestas.docx", doc_listening(mat["titulo"], mat["desc"], mat["frases"], mat["dictado"], True))
    st.download_button(
        "⬇ Descargar todo (audios MP3 + hoja + clave)", paquete.getvalue(), "listening.zip", "application/zip",
        type="primary", use_container_width=True, key="pl_zip",
    )
    for i, (f, audio) in enumerate(zip(mat["frases"], mat["audios"]), 1):
        with st.container(border=True):
            st.markdown(f"**{i}. {f['zh']}**  \n{f['pinyin']} — {f['es']}")
            if audio:
                st.audio(audio, format="audio/mp3")


def _tab_vocabulario(tutor, vocab, frases, raiz: Path, leer_secreto, recargar):
    S = st.session_state
    st.caption("Agrega un vocabulario nuevo o corrige uno existente, con sus frases de entrenamiento.")
    if aviso := S.pop("pv_aviso", None):
        st.success(aviso)

    niveles = sorted(vocab)
    a, b = st.columns(2)
    # Las claves de los selectores cambian cuando cambia el curso, para que tras guardar
    # se redibujen apuntando a la unidad recién guardada.
    ult_nivel, ult_unidad = S.get("pv_ultimo", (None, None))
    op_niveles = niveles + [niveles[-1] + 1]
    nivel = a.selectbox(
        "Nivel", op_niveles, key=f"pv_nivel_{len(niveles)}",
        index=op_niveles.index(ult_nivel) if ult_nivel in op_niveles else 0,
        format_func=lambda n: f"Chino {n}" if n in vocab else f"➕ Nuevo nivel: Chino {n}",
    )
    existentes = sorted(vocab.get(nivel, {}))
    nueva = (existentes[-1] + 1) if existentes else 1
    op_unidades = [nueva] + existentes
    unidad = b.selectbox(
        "Vocabulario", op_unidades, key=f"pv_unidad_{nivel}_{len(existentes)}",
        index=op_unidades.index(ult_unidad) if nivel == ult_nivel and ult_unidad in op_unidades else 0,
        format_func=lambda u: f"➕ Nuevo: Vocabulario {u}" if u == nueva else f"Editar Vocabulario {u}",
    )
    k_pal, k_fra = f"pv_pal_{nivel}_{unidad}", f"pv_fra_{nivel}_{unidad}"
    if k_fra in S.get("pv_sugeridas", {}):  # frases propuestas por la IA en la pasada anterior
        S[k_fra] = S.pv_sugeridas.pop(k_fra)

    S.setdefault(k_pal, "， ".join(vocab.get(nivel, {}).get(unidad, [])))
    S.setdefault(k_fra, "\n".join(frases.get(nivel, {}).get(unidad, [])))
    txt_pal = st.text_area(
        "Palabras (separadas por comas o una por línea)", key=k_pal, height=110, placeholder="天气， 冷， 热， 下雨"
    )
    txt_fra = st.text_area(
        "Frases de entrenamiento (una por línea)", key=k_fra, height=230, placeholder="今天天气很冷。\n明天热吗？"
    )
    palabras = core.separar_palabras(txt_pal)
    lineas = [l.strip() for l in txt_fra.splitlines() if l.strip()]

    # Vocabulario que el estudiante conocería al llegar a esta unidad (incluida la nueva)
    vocab_tmp = {n: dict(us) for n, us in vocab.items()}
    vocab_tmp.setdefault(nivel, {})[unidad] = palabras or ["?"]
    frases_tmp = {n: dict(us) for n, us in frases.items()}
    frases_tmp.setdefault(nivel, {})[unidad] = []
    ctx = core.construir_contexto(vocab_tmp, frases_tmp, nivel, [unidad])

    if st.button("✨ Sugerir frases con IA", disabled=not palabras, key="pv_sugerir"):
        with st.spinner("Escribiendo frases con este vocabulario…"):
            try:
                sugeridas = tutor.sugerir_frases(ctx)
            except Exception as err:
                st.error(f"No se pudieron generar frases ({type(err).__name__}). Intenta de nuevo.")
                sugeridas = []
        if sugeridas:
            S.setdefault("pv_sugeridas", {})[k_fra] = "\n".join(lineas + [f for f in sugeridas if f not in lineas])
            st.rerun()

    if palabras:
        st.caption(f"{len(palabras)} palabras · {len(lineas)} frases")
    sin_usar = [p for p in palabras if not any(p in l for l in lineas)]
    if lineas and sin_usar:
        st.info("Palabras que aún no aparecen en ninguna frase: " + "， ".join(sin_usar))
    extra = core.fuera_de_vocabulario("".join(lineas), ctx)
    if extra:
        st.warning(
            "Estos caracteres aparecen en las frases pero no están en ningún vocabulario hasta esta unidad: "
            + " ".join(extra) + ". Puedes dejarlos (la app los aceptará en esta unidad) o agregarlos a las palabras."
        )

    if not st.button("Guardar vocabulario", type="primary", disabled=not (palabras and lineas), key="pv_guardar"):
        if not (palabras and lineas):
            st.caption("Escribe las palabras y al menos una frase para poder guardar.")
        _descargas_pendientes()
        return

    nuevo_vocab = core.guardar_unidad((raiz / ARCHIVO_VOCAB).read_text(encoding="utf-8"), nivel, unidad, ["， ".join(palabras)])
    nuevas_frases = core.guardar_unidad((raiz / ARCHIVO_FRASES).read_text(encoding="utf-8"), nivel, unidad, lineas)
    archivos = {ARCHIVO_VOCAB: nuevo_vocab, ARCHIVO_FRASES: nuevas_frases}
    token = leer_secreto("GITHUB_TOKEN")
    guardado_github = False
    if token:
        try:
            subir_a_github(
                token, leer_secreto("GITHUB_REPO", "kaezar22/mandarin_asistant"), leer_secreto("GITHUB_BRANCH", "main"),
                archivos, f"Chino {nivel}: vocabulario {unidad} (desde la zona del profesor)",
            )
            guardado_github = True
        except Exception as err:
            st.error(f"No se pudo guardar en GitHub ({type(err).__name__}). Revisa el token; abajo puedes descargar los archivos.")
    try:  # efecto inmediato en esta sesión de la app
        for nombre, texto in archivos.items():
            (raiz / nombre).write_text(texto, encoding="utf-8")
        recargar()
    except OSError:
        pass
    S.pv_ultimo = (nivel, unidad)
    if guardado_github:
        S.pv_aviso = f"Guardado: Chino {nivel} · Vocabulario {unidad}. Quedó en GitHub y la app se actualiza en un par de minutos."
        S.pop("pv_descargas", None)
    else:
        S.pv_descargas = archivos
        S.pv_aviso = (
            f"Chino {nivel} · Vocabulario {unidad} ya se puede usar en esta sesión, pero es temporal: "
            "para que quede fijo, descarga los dos archivos de abajo y súbelos al repositorio de GitHub."
        )
    st.rerun()


def _descargas_pendientes():
    archivos = st.session_state.get("pv_descargas")
    if not archivos:
        return
    a, b = st.columns(2)
    for col, (nombre, texto) in zip((a, b), archivos.items()):
        col.download_button(f"⬇ {nombre}", texto.encode("utf-8"), nombre, "text/plain", use_container_width=True, key=f"pv_d_{nombre}")


def render(tutor, vocab, frases, raiz: Path, leer_secreto, sintetizar, recargar):
    st.subheader("🔒 Zona del profesor")
    if not _acceso(leer_secreto("CLAVE_PROFESOR", "2209")):
        return
    st.markdown("**Material para:**")
    ctx, desc = _material_para(vocab, frases)
    t1, t2, t3, t4 = st.tabs(["📝 Examen", "✍️ Escritura", "🎧 Listening", "➕ Vocabulario"])
    with t1:
        _tab_examen(tutor, ctx, desc)
    with t2:
        _tab_escritura(tutor, ctx, desc)
    with t3:
        _tab_listening(tutor, ctx, desc, sintetizar)
    with t4:
        _tab_vocabulario(tutor, vocab, frases, raiz, leer_secreto, recargar)
